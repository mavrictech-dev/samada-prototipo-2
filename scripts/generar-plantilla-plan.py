"""Genera el borrador de plantilla del plan; no implementa el importador del CMS."""
from pathlib import Path
import csv
import json
import zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'plantillas'
OUT.mkdir(parents=True, exist_ok=True)
GREEN = '558B2F'
DARK = '24321D'
INPUT = 'F3F8ED'
BLUE = '1F4E78'
ROWS = 1500

# Encabezado, obligatorio al crear, tipo/regla, vacío al crear, ejemplo
FIELDS = [
 ('sku', 'SI', 'Texto único. Conservar ceros iniciales. Letras, números, punto, guion y guion bajo.', '', '000123'),
 ('nombre', 'SI', 'Texto comercial.', '', 'Shampoo Nutritivo 250 ml'),
 ('marca_codigo', 'SI', 'Código válido de Marcas.', '', 'BOTANICA'),
 ('categoria_codigo', 'SI', 'Código de Categorias habilitado para esa marca.', '', 'CAPILAR'),
 ('subcategoria_codigo', 'NO', 'Código de Subcategorias válido para categoría y marca.', 'Sin subcategoría', 'SHAMPOO'),
 ('tamano_codigo', 'SI', 'Código de Tamanos habilitado para marca. SIN-TAMANO si no aplica.', '', 'VOL-250ML'),
 ('codigo_grupo', 'NO', 'Texto. Agrupa varios SKU de una misma familia.', 'Sin grupo', 'SHAMPOO-NUTRITIVO'),
 ('presentacion', 'NO', 'Texto comercial adicional.', 'Vacío', 'Botella'),
 ('precio_mayorista', 'SI', 'Decimal >= 0. Punto decimal. Sin S/ ni separador de miles.', '', 19.9),
 ('precio_sugerido', 'NO', 'Decimal >= 0, opcional.', 'Vacío', 28.9),
 ('stock', 'NO', 'Entero >= 0. Cantidad absoluta, no incremento.', '0', 100),
 ('pedido_minimo', 'NO', 'Entero >= 1.', '1', 6),
 ('unidades_caja', 'NO', 'Entero >= 1.', '1', 12),
 ('descripcion', 'NO', 'Texto sin fórmulas.', 'Vacío', 'Limpieza suave para la rutina diaria.'),
 ('uso', 'NO', 'Texto sin fórmulas.', 'Vacío', 'Uso capilar'),
 ('publicado', 'NO', 'SI o NO. La publicación requiere permiso.', 'NO', 'NO'),
 ('destacado', 'NO', 'SI o NO.', 'NO', 'NO'),
 ('orden', 'NO', 'Entero >= 0. Vacío para orden automático.', 'Automático', 10),
 ('moneda', 'NO', 'Código de moneda configurado. Inicialmente PEN.', 'PEN', 'PEN'),
 ('imagen_archivo', 'NO', 'Nombre exacto de archivo del lote. Si vacío, buscar por SKU.', 'Buscar por SKU', '000123.jpg'),
]
HEADERS = [f[0] for f in FIELDS]
wb = Workbook()
wb.remove(wb.active)

def style_sheet(ws, widths=None, input_sheet=False):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30
    for row in ws:
        for cell in row:
            cell.font = Font(name='Arial', size=10, color='24321D')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor=DARK)
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for i in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[i-1] if widths and i <= len(widths) else 24
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'

guide = wb.create_sheet('Instrucciones')
guide.append(['SAMADA · Plantilla de catálogo v1 — borrador del plan', 'Cómo usar'])
instructions = [
 ('Estado', 'Borrador de campos para la implementación. El sistema actual todavía no importa este archivo.'),
 ('Celdas editables', 'Rellena las filas verdes de Productos. Las columnas con encabezado verde son obligatorias para crear SKU nuevos.'),
 ('Una fila', 'Una presentación vendible = un SKU. Cambiar de tamaño implica otro SKU. codigo_grupo reúne variantes.'),
 ('Ejemplos', 'La hoja Ejemplos contiene datos demostrativos. No es parte de la carga. Productos empieza vacío.'),
 ('Campos obligatorios', 'sku, nombre, marca_codigo, categoria_codigo, tamano_codigo y precio_mayorista.'),
 ('Códigos', 'Usa las listas y revisa MarcaCategoria, MarcaSubcategoria y MarcaTamano. Son ejemplos hasta conectar el generador al CMS.'),
 ('SKU como texto', 'sku y todos los códigos son texto. No perder ceros iniciales: 000123 es distinto de 123.'),
 ('Tamaños', '250 ml, 500 ml y 1 L se ordenan por volumen real. No comparar masa con volumen. SIN-TAMANO cuando no aplica.'),
 ('Números', 'Escribir 19.90, sin S/ ni separadores de miles. Stock entero no negativo. Pedido mínimo y unidades/caja >= 1.'),
 ('Valores iniciales', 'Al crear: stock 0; pedido_minimo 1; unidades_caja 1; publicado NO; destacado NO; moneda PEN; orden automático.'),
 ('Actualizar', 'Modo actualizar: sku obligatorio; columnas ausentes o vacías conservan el valor anterior. 0 y NO sí actualizan.'),
 ('Vaciar explícitamente', '__BORRAR__ se admitirá solo en campos opcionales autorizados. No lo uses en campos obligatorios.'),
 ('Crear/actualizar', 'Un SKU existente se actualiza. Un SKU nuevo debe tener todos los campos obligatorios. Ausentes del archivo no se eliminan.'),
 ('Imágenes', 'Carga por separado imágenes con nombre SKU.jpg, SKU.png o SKU.webp. El motor convertirá a WebP y revisará duplicados.'),
 ('Galería propuesta', 'Si se habilita galería: SKU__01.jpg, SKU__02.png. Un conflicto requiere revisión; no se asigna por similitud.'),
 ('CSV', 'UTF-8 con encabezados idénticos, coma delimitadora y punto decimal en la plantilla. Cuidado: CSV no conserva tipos ni listas.'),
 ('Filas preparadas', 'Productos incluye formato y listas para 1.500 filas. El importador deberá ignorar filas vacías.'),
 ('Fórmulas', 'Pega valores. No se aceptarán fórmulas ni macros en los campos de importación.'),
 ('Versión de esquema', 'samada-catalogo-v1. En el sistema, la plantilla se generará con los catálogos actuales y metadatos de versión.'),
 ('Origen', 'Campos derivados de los requisitos del negocio y decisiones propuestas del plan del 19/10/2026. Ejemplos sin valor comercial real.'),
]
for row in instructions:
    guide.append(row)
style_sheet(guide, [42, 115])
for r in range(2, guide.max_row + 1):
    guide.row_dimensions[r].height = 44

products = wb.create_sheet('Productos')
products.append(HEADERS)
for r in range(2, ROWS + 2):
    for c, field in enumerate(FIELDS, 1):
        cell = products.cell(r, c)
        cell.fill = PatternFill('solid', fgColor=INPUT if field[1] == 'SI' else 'FBFCF8')
        cell.font = Font(name='Arial', size=10, color=BLUE)
        cell.alignment = Alignment(vertical='top')
        if field[0] in {'precio_mayorista','precio_sugerido'}:
            cell.number_format = '0.00'
        elif field[0] in {'stock','pedido_minimo','unidades_caja','orden'}:
            cell.number_format = '0'
        else:
            cell.number_format = '@'
style_sheet(products)
for c, field in enumerate(FIELDS, 1):
    cell = products.cell(1, c)
    cell.comment = Comment(f"{'OBLIGATORIO al crear' if field[1] == 'SI' else 'Opcional'}\n{field[2]}\nVacío al crear: {field[3] or 'No permitido'}\nEjemplo: {field[4]}", 'Samada')
    if field[1] == 'SI':
        cell.fill = PatternFill('solid', fgColor=GREEN)
    products.column_dimensions[get_column_letter(c)].width = 42 if field[0] in {'nombre','descripcion','uso'} else 23
table = Table(displayName='ProductosCarga', ref=f'A1:T{ROWS + 1}')
table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium4', showRowStripes=False)
products.add_table(table)

examples = wb.create_sheet('Ejemplos')
examples.append(HEADERS)
example1 = [f[4] for f in FIELDS]
examples.append(example1)
example2 = list(example1)
example2[0] = '000124'
example2[1] = 'Shampoo Nutritivo 500 ml'
example2[5] = 'VOL-500ML'
example2[8], example2[9], example2[10], example2[17], example2[19] = 32.9, 45.9, 60, 20, '000124.jpg'
examples.append(example2)
style_sheet(examples)
for row in examples.iter_rows(min_row=2):
    row[0].number_format = '@'
    row[8].number_format = row[9].number_format = '0.00'

fields = wb.create_sheet('Campos')
fields.append(['campo', 'obligatorio_crear', 'regla', 'valor_si_vacio_al_crear', 'ejemplo'])
for f in FIELDS:
    fields.append(f)
style_sheet(fields, [26, 23, 85, 30, 43])

references = {
 'Marcas': (['codigo','nombre','estado'], [('BOTANICA','Botánica','ACTIVO'),('ESSENTIA','Essentia','ACTIVO'),('SKINLAB','Skinlab','ACTIVO')]),
 'Categorias': (['codigo','nombre','orden'], [('CAPILAR','Cuidado capilar',10),('FACIAL','Cuidado facial',20),('CORPORAL','Cuidado corporal',30)]),
 'Subcategorias': (['codigo','nombre','categoria_codigo'], [('SHAMPOO','Shampoo','CAPILAR'),('TRATAMIENTOS','Tratamientos','CAPILAR'),('SERUMS','Sérums','FACIAL'),('CREMAS','Cremas','CORPORAL')]),
 'Tamanos': (['codigo','etiqueta','tipo','valor','unidad','valor_base','unidad_base','orden'], [('VOL-250ML','250 ml','VOLUMEN',250,'ml',250,'ml',10),('VOL-500ML','500 ml','VOLUMEN',500,'ml',500,'ml',20),('VOL-1L','1 L','VOLUMEN',1,'L',1000,'ml',30),('MAS-250G','250 g','MASA',250,'g',250,'g',10),('TALLA-S','S','TEXTO',None,None,None,None,10),('TALLA-M','M','TEXTO',None,None,None,None,20),('TALLA-L','L','TEXTO',None,None,None,None,30),('SIN-TAMANO','Sin tamaño','TEXTO',None,None,None,None,999)]),
 'MarcaCategoria': (['marca_codigo','categoria_codigo','orden'], [('BOTANICA','CAPILAR',10),('BOTANICA','CORPORAL',20),('ESSENTIA','CAPILAR',10),('ESSENTIA','CORPORAL',20),('SKINLAB','FACIAL',10)]),
 'MarcaSubcategoria': (['marca_codigo','subcategoria_codigo'], [('BOTANICA','SHAMPOO'),('BOTANICA','TRATAMIENTOS'),('BOTANICA','CREMAS'),('ESSENTIA','SHAMPOO'),('ESSENTIA','TRATAMIENTOS'),('ESSENTIA','CREMAS'),('SKINLAB','SERUMS')]),
 'MarcaTamano': (['marca_codigo','tamano_codigo'], [(brand, size) for brand in ['BOTANICA','ESSENTIA','SKINLAB'] for size in ['VOL-250ML','VOL-500ML','VOL-1L','MAS-250G','TALLA-S','TALLA-M','TALLA-L','SIN-TAMANO']]),
}
for name, (header, rows) in references.items():
    ws = wb.create_sheet(name)
    ws.append(header)
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    ws['A1'].comment = Comment('Datos de referencia demostrativos. El CMS generará esta hoja con su configuración vigente.', 'Samada')

def list_validation(column, formula):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True)
    dv.errorTitle, dv.error = 'Código no válido', 'Elige un valor de las hojas de referencia.'
    dv.showErrorMessage = True
    dv.errorStyle = 'stop'
    products.add_data_validation(dv)
    dv.add(f'{column}2:{column}{ROWS+1}')

for column, name in [('C','Marcas'),('D','Categorias'),('E','Subcategorias'),('F','Tamanos')]:
    defined = DefinedName(f'Lista{name}', attr_text=f"'{name}'!$A$2:$A${wb[name].max_row}")
    wb.defined_names.add(defined)
    list_validation(column, f'=Lista{name}')
list_validation('P', '"SI,NO"')
list_validation('Q', '"SI,NO"')
list_validation('S', '"PEN"')
for column, kind, lower in [('I','decimal',0),('J','decimal',0),('K','whole',0),('L','whole',1),('M','whole',1),('R','whole',0)]:
    dv = DataValidation(type=kind, operator='greaterThanOrEqual', formula1=str(lower), allow_blank=True)
    dv.showErrorMessage, dv.errorTitle, dv.error = True, 'Número no válido', f'Ingresa un número {">="} {lower}.'
    products.add_data_validation(dv)
    dv.add(f'{column}2:{column}{ROWS+1}')

wb.properties.title = 'Samada · Plantilla de catálogo v1'
wb.properties.subject = 'Borrador del plan: SKU, marca, categoría, tamaño e importación masiva'
wb.properties.creator = 'Samada'
wb.active = 0
xlsx = OUT / 'Samada-Plantilla-Catalogo-v1.xlsx'
wb.save(xlsx)
for filename, rows in [('Samada-Plantilla-Catalogo-v1.csv', []), ('Samada-Ejemplos-Catalogo-v1.csv', [example1,example2])]:
    with (OUT / filename).open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(rows)

# Verificación estructural: encabezados, filas vacías, SKU de ejemplo y listas.
check = load_workbook(xlsx)
assert [c.value for c in check['Productos'][1]] == HEADERS
assert all(c.value is None for row in check['Productos'].iter_rows(min_row=2) for c in row)
assert check['Ejemplos']['A2'].value == '000123'
assert check['Ejemplos']['A2'].number_format == '@'
assert len(check['Productos'].data_validations.dataValidation) == 13
assert len(check.sheetnames) == 11
assert not any(c.data_type == 'f' for ws in check for row in ws for c in row)
with zipfile.ZipFile(xlsx) as archive:
    assert archive.testzip() is None
summary = {'xlsx': str(xlsx), 'hojas': check.sheetnames, 'columnas': len(HEADERS), 'filas_preparadas': ROWS, 'formulas': 0, 'verificacion': 'OK'}
(OUT / 'verificacion-plantilla.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(summary, ensure_ascii=True))
